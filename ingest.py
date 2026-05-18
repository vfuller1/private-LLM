"""
ingest.py — Read documents from docs/, chunk them, embed each chunk
with a local model, and store the embeddings in a local Chroma database.

Run this whenever you add or change files in docs/:

    python ingest.py
"""

from __future__ import annotations

import sys
from pathlib import Path

import chromadb
import ollama
from docx import Document
from pypdf import PdfReader

import config


# File extensions this script knows how to read.
SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".md"}


# ---------------------------------------------------------------------------
# Document loading
# ---------------------------------------------------------------------------

def load_text_file(path: Path) -> str:
    """Read a plain-text file."""
    return path.read_text(encoding="utf-8", errors="ignore")


def load_pdf(path: Path) -> str:
    """Extract text from every page of a PDF."""
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def load_docx(path: Path) -> str:
    """
    Extract text from a Word .docx file.

    We pull text from paragraphs and from any tables in the document.
    Old .doc files (Word 97-2003 format) are NOT supported - convert
    them to .docx in Word first (File > Save As > .docx).
    """
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                parts.append(row_text)
    return "\n\n".join(parts)


def load_document(path: Path) -> str:
    """Dispatch to the right loader based on file extension."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".docx":
        return load_docx(path)
    if suffix in {".txt", ".md"}:
        return load_text_file(path)
    raise ValueError(f"Unsupported file type: {suffix}")


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

def chunk_text(text, size, overlap):
    """
    Split text into overlapping chunks of roughly `size` characters.

    We use a simple sliding window. Real-world systems often split on
    sentence or paragraph boundaries for better results, but this is
    plenty for a first project.
    """
    text = text.strip()
    if not text:
        return []

    chunks = []
    start = 0
    step = max(1, size - overlap)

    while start < len(text):
        end = min(start + size, len(text))
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == len(text):
            break
        start += step

    return chunks


# ---------------------------------------------------------------------------
# Embedding
# ---------------------------------------------------------------------------

def embed_one(text):
    """
    Ask Ollama to turn a string into a vector using the embedding model.
    The vector is a list of floats representing the text's meaning.
    """
    response = ollama.embeddings(model=config.EMBED_MODEL, prompt=text)
    return response["embedding"]


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    docs_dir = config.DOCS_DIR
    if not docs_dir.exists():
        print(f"No docs folder found at {docs_dir}.")
        print("Create it and drop some .pdf/.docx/.txt/.md files inside, then re-run.")
        sys.exit(1)

    client = chromadb.PersistentClient(path=str(config.DB_DIR))

    try:
        client.delete_collection(config.COLLECTION_NAME)
    except Exception:
        pass
    collection = client.create_collection(config.COLLECTION_NAME)

    files = sorted(
        p for p in docs_dir.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS
    )
    if not files:
        exts = ", ".join(sorted(SUPPORTED_EXTS))
        print(f"No supported files in {docs_dir}. Add files of type: {exts}")
        sys.exit(1)

    print(f"Loading documents from {docs_dir}/")

    all_ids = []
    all_documents = []
    all_metadatas = []

    for path in files:
        try:
            raw_text = load_document(path)
        except Exception as exc:
            print(f"  ! Skipping {path.name}: {exc}")
            continue

        chunks = chunk_text(raw_text, config.CHUNK_SIZE, config.CHUNK_OVERLAP)
        print(f"  - {path.name}  ({len(chunks)} chunks)")

        for i, chunk in enumerate(chunks):
            all_ids.append(f"{path.name}::chunk-{i}")
            all_documents.append(chunk)
            all_metadatas.append({"source": path.name, "chunk": i})

    if not all_documents:
        print("Nothing to embed - exiting.")
        sys.exit(1)

    print(f"Embedding {len(all_documents)} chunks with {config.EMBED_MODEL}...")

    embeddings = []
    for idx, chunk in enumerate(all_documents, start=1):
        embeddings.append(embed_one(chunk))
        if idx % 10 == 0 or idx == len(all_documents):
            print(f"  {idx}/{len(all_documents)} embedded")

    collection.add(
        ids=all_ids,
        documents=all_documents,
        metadatas=all_metadatas,
        embeddings=embeddings,
    )

    print(f"Done. Stored {len(all_documents)} chunks in {config.DB_DIR}/")
    print("Next step: run `python chat.py` or `python quiz.py`")


if __name__ == "__main__":
    main()
